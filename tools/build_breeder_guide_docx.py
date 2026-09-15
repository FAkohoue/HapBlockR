#!/usr/bin/env python3
"""Build the editable HapBlockR Breeder's Guide Word edition.

The Markdown file in inst/guide is the single editorial source for both the
tagged PDF and this DOCX.  The builder uses native Word headings, lists,
tables, hyperlinks, page numbers and language metadata so that the document
remains editable and accessible.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "inst" / "guide" / "HapBlockR_Breeder_Guide.md"
OUTPUT = ROOT / "inst" / "extdata" / "HapBlockR_Breeder_Guide.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
GREEN = "3D7D44"
PALE_BLUE = "E8EEF5"
LIGHT_GREY = "F2F4F7"
BODY_COLOUR = RGBColor(23, 32, 51)
# A4 width minus the 16 mm left and right margins, expressed in twentieths
# of a point for native Word table geometry.
USABLE_WIDTH_DXA = round((210 - 32) / 25.4 * 1440)


def set_cell_margins(cell, top=60, start=120, bottom=60, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def new_numbering_id(document):
    """Create a fresh level-zero sequence based on Word's List Number style."""
    numbering = document.part.numbering_part.element
    style_num_id = document.styles["List Number"].element.pPr.numPr.numId.val
    source_num = next(
        node
        for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == style_num_id
    )
    abstract_id = source_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    num_id = max(existing, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.extend([abstract, override])
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def add_hyperlink(paragraph, text, url):
    relationship = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend([colour, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([run_properties, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(
    r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\)|https?://\S+)"
)


def add_inline(paragraph, text, *, base_size=None):
    """Add a small, deterministic subset of Markdown inline formatting."""
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            if base_size:
                run.font.size = Pt(base_size)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt((base_size or 11) - 0.5)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), LIGHT_GREY)
            run._r.get_or_add_rPr().append(shade)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            if base_size:
                run.font.size = Pt(base_size)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            if base_size:
                run.font.size = Pt(base_size)
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        else:
            url = token.rstrip(".,;")
            add_hyperlink(paragraph, url, url)
            suffix = token[len(url) :]
            if suffix:
                paragraph.add_run(suffix)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        if base_size:
            run.font.size = Pt(base_size)


def set_language(style, language="en-GB"):
    r_pr = style.element.get_or_add_rPr()
    language_node = r_pr.find(qn("w:lang"))
    if language_node is None:
        language_node = OxmlElement("w:lang")
        r_pr.append(language_node)
    language_node.set(qn("w:val"), language)
    language_node.set(qn("w:eastAsia"), language)
    language_node.set(qn("w:bidi"), language)


def set_style_font(style, name):
    style.font.name = name
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


def configure_list_numbering(document):
    """Apply the compact-reference list geometry to every level-zero list."""
    numbering = document.part.numbering_part.element
    for abstract in numbering.findall(qn("w:abstractNum")):
        level = next(
            (
                node
                for node in abstract.findall(qn("w:lvl"))
                if node.get(qn("w:ilvl")) == "0"
            ),
            None,
        )
        if level is None:
            continue
        justification = level.find(qn("w:lvlJc"))
        if justification is None:
            justification = OxmlElement("w:lvlJc")
            level.append(justification)
        justification.set(qn("w:val"), "left")
        p_pr = level.find(qn("w:pPr"))
        if p_pr is None:
            p_pr = OxmlElement("w:pPr")
            level.append(p_pr)
        tabs = p_pr.find(qn("w:tabs"))
        if tabs is None:
            tabs = OxmlElement("w:tabs")
            p_pr.append(tabs)
        for old_tab in list(tabs):
            tabs.remove(old_tab)
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "540")
        tabs.append(tab)
        indent = p_pr.find(qn("w:ind"))
        if indent is None:
            indent = OxmlElement("w:ind")
            p_pr.append(indent)
        indent.set(qn("w:left"), "540")
        indent.set(qn("w:hanging"), "270")


def configure_styles(document):
    styles = document.styles
    normal = styles["Normal"]
    set_style_font(normal, "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = BODY_COLOUR
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, colour, before, after) in heading_specs.items():
        style = styles[name]
        set_style_font(style, "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(colour)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = styles["Title"]
    set_style_font(title, "Calibri")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    title.paragraph_format.space_after = Pt(14)

    subtitle = styles["Subtitle"]
    set_style_font(subtitle, "Calibri")
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(GREEN)
    subtitle.paragraph_format.space_after = Pt(18)

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    set_style_font(code_style, "Consolas")
    code_style.font.size = Pt(9)
    code_style.paragraph_format.left_indent = Inches(0.25)
    code_style.paragraph_format.right_indent = Inches(0.15)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(6)
    code_style.paragraph_format.keep_together = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    for style in styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            set_language(style)
    configure_list_numbering(document)


def configure_document(document):
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(17)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(16)
    section.right_margin = Mm(16)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("HapBlockR Breeder's Guide | Edition 6")
    run.font.name = "Calibri"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("HBR-GUIDE-001  |  ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(89, 89, 89)
    add_page_field(paragraph)

    document.core_properties.title = "The HapBlockR Breeder's Guide"
    document.core_properties.subject = (
        "Haplotype-aware parent selection and breeding decisions"
    )
    document.core_properties.author = "Félicien Akohoue"
    document.core_properties.keywords = (
        "HapBlockR, breeding, haplotypes, parent selection, genomic selection"
    )
    document.core_properties.comments = (
        "Editable Word edition generated from the version-controlled guide source."
    )


def table_widths(rows):
    n_cols = max(len(row) for row in rows)
    scores = []
    for col in range(n_cols):
        values = [row[col] if col < len(row) else "" for row in rows]
        longest = min(max((len(value) for value in values), default=1), 55)
        scores.append(max(7.0, longest ** 0.72))
    total = sum(scores)
    widths = [max(560, round(USABLE_WIDTH_DXA * score / total)) for score in scores]
    difference = USABLE_WIDTH_DXA - sum(widths)
    widths[-1] += difference
    return widths


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_pr = table._tbl.tblPr
    table_width = table_pr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_pr.append(table_width)
    table_width.set(qn("w:w"), str(USABLE_WIDTH_DXA))
    table_width.set(qn("w:type"), "dxa")

    indent = table_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        table_pr.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")

    old_grid = table._tbl.tblGrid
    new_grid = OxmlElement("w:tblGrid")
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        new_grid.append(column)
    table._tbl.replace(old_grid, new_grid)

    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            width = widths[index]
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_width = tc_pr.find(qn("w:tcW"))
            if tc_width is None:
                tc_width = OxmlElement("w:tcW")
                tc_pr.append(tc_width)
            tc_width.set(qn("w:w"), str(width))
            tc_width.set(qn("w:type"), "dxa")


def parse_table_line(line):
    return [part.strip() for part in line.strip().strip("|").split("|")]


def is_separator_row(row):
    return all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in row)


def add_table(document, raw_lines):
    rows = [parse_table_line(line) for line in raw_lines]
    if len(rows) > 1 and is_separator_row(rows[1]):
        rows.pop(1)
    n_cols = max(len(row) for row in rows)
    rows = [row + [""] * (n_cols - len(row)) for row in rows]
    table = document.add_table(rows=len(rows), cols=n_cols)
    table.style = "Table Grid"
    widths = table_widths(rows)
    set_table_geometry(table, widths)

    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_inline(paragraph, value, base_size=8.0)
            if len(rows) <= 10 and row_index < len(rows) - 1:
                paragraph.paragraph_format.keep_with_next = True
            if row_index == 0:
                shade_cell(cell, PALE_BLUE)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        if row_index == 0:
            set_repeat_table_header(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_code_block(document, lines):
    paragraph = document.add_paragraph(style="Code Block")
    add_inline(paragraph, "\n".join(lines), base_size=9)
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_GREY)
    p_pr.append(shading)


def add_body_paragraph(document, text, cover=False):
    paragraph = document.add_paragraph()
    if cover:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(5)
    add_inline(paragraph, text)
    return paragraph


def build():
    text = SOURCE.read_text(encoding="utf-8")
    lines = text.splitlines()
    document = Document()
    configure_styles(document)
    configure_document(document)

    paragraph_buffer = []
    in_code = False
    code_buffer = []
    before_first_break = True
    title_seen = False
    subtitle_seen = False
    active_numbering_id = None
    last_list_paragraph = None

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            joined = " ".join(part.strip() for part in paragraph_buffer).strip()
            paragraph = add_body_paragraph(
                document, joined, cover=before_first_break
            )
            if re.match(r"^Edition \d+(?:,|:)", joined):
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.15
            paragraph_buffer = []

    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_paragraph()
            active_numbering_id = None
            last_list_paragraph = None
            if in_code:
                add_code_block(document, code_buffer)
                code_buffer = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_buffer.append(line)
            index += 1
            continue

        if stripped == "<!-- pagebreak -->":
            flush_paragraph()
            active_numbering_id = None
            last_list_paragraph = None
            document.add_page_break()
            before_first_break = False
            index += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            active_numbering_id = None
            last_list_paragraph = None
            level = len(heading.group(1))
            content = heading.group(2)
            if level == 1 and not title_seen:
                paragraph = document.add_paragraph(style="Title")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, content)
                title_seen = True
            elif level == 2 and before_first_break and not subtitle_seen:
                paragraph = document.add_paragraph(style="Subtitle")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(paragraph, content)
                subtitle_seen = True
            else:
                word_level = min(3, max(1, level - 1))
                paragraph = document.add_paragraph(style=f"Heading {word_level}")
                add_inline(paragraph, content)
            index += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            active_numbering_id = None
            last_list_paragraph = None
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
            continue

        bullet = re.match(r"^\s*-\s+(.+)$", line)
        numbered = re.match(r"^\s*\d+\.\s+(.+)$", line)
        if bullet or numbered:
            flush_paragraph()
            paragraph = document.add_paragraph(
                style="List Bullet" if bullet else "List Number"
            )
            if numbered:
                if active_numbering_id is None:
                    active_numbering_id = new_numbering_id(document)
                apply_numbering(paragraph, active_numbering_id)
            else:
                active_numbering_id = None
            add_inline(paragraph, (bullet or numbered).group(1))
            last_list_paragraph = paragraph
            index += 1
            continue

        if re.match(r"^\s{2,}\S", line) and last_list_paragraph is not None:
            last_list_paragraph.add_run(" ")
            add_inline(last_list_paragraph, stripped)
            index += 1
            continue

        if not stripped:
            flush_paragraph()
            active_numbering_id = None
            last_list_paragraph = None
            index += 1
            continue

        if before_first_break and re.match(
            r"^(Document ID|Edition|Compatible package version|Release date|Author|Source):",
            stripped,
        ):
            flush_paragraph()
            active_numbering_id = None
            last_list_paragraph = None
            add_body_paragraph(document, stripped.rstrip(), cover=True)
            index += 1
            continue

        active_numbering_id = None
        last_list_paragraph = None
        paragraph_buffer.append(line)
        index += 1

    flush_paragraph()
    if in_code:
        add_code_block(document, code_buffer)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Built {OUTPUT}")


if __name__ == "__main__":
    build()
